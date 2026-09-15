"""Document extraction: PDF, DOCX, HWPX, HWP 5.0 records, and intake integration."""
from __future__ import annotations

import io
import struct
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest

from claim_agent.models.request import MaterialBundle, RunRequest
from claim_agent.sources import extract as X
from claim_agent.tui_support import validate_attachment


def _pdf_with_text(lines: list[str]) -> bytes:
    """Minimal single-page PDF with a Helvetica text stream (ASCII only)."""
    content = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_pdf_text_layer_and_scanned_rejection():
    r = X.extract_pdf(_pdf_with_text(["Claim 1. A holder comprising a base.", "Claim 2. The holder of claim 1."]))
    assert r.format == "pdf" and r.pages == 1 and "claim 1" in r.text.lower() and "[페이지 1]" in r.text
    with pytest.raises(X.ExtractionError, match="스캔본"):
        X.extract_pdf(_pdf_with_text(["x"]))


def test_docx_paragraphs_and_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("【청구항 1】")
    d.add_paragraph("베이스; 및 홀더 본체를 포함하는 케이블 클립 홀더.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "구성", "근거"
    buf = io.BytesIO()
    d.save(buf)
    r = X.extract_docx(buf.getvalue())
    assert r.format == "docx" and "【청구항 1】\n베이스; 및 홀더 본체" in r.text and "구성\t근거" in r.text


def _hwpx(paragraphs: list[str]) -> bytes:
    ns = "http://www.hancom.co.kr/hwpml/2011/paragraph"
    body = "".join(f'<hp:p><hp:run><hp:t>{p}</hp:t></hp:run></hp:p>' for p in paragraphs)
    xml = f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="{ns}">{body}</hs:sec>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/hwp+zip")
        zf.writestr("Contents/section0.xml", xml)
    return buf.getvalue()


def test_hwpx_paragraphs():
    r = X.extract_hwpx(_hwpx(["【청구항 1】", "제1항에 있어서, 탄성 리브를 더 포함하는 홀더."]))
    assert r.format == "hwpx" and r.text == "【청구항 1】\n제1항에 있어서, 탄성 리브를 더 포함하는 홀더."
    with pytest.raises(X.ExtractionError):
        X.extract_hwpx(b"not a zip")


def _hwp_section(paragraphs: list[str]) -> bytes:
    out = b""
    for p in paragraphs:
        # 11 = drawing control and 9 = tab are 8-wchar controls in HWP 5.0 paragraph text
        payload = b"".join(struct.pack("<H", 11) + b"\0" * 14 if ch == "\ufffc" else struct.pack("<H", 9) + b"\0" * 14 if ch == "\t" else ch.encode("utf-16-le") for ch in p)
        payload += struct.pack("<H", 13)
        header = (X._HWPTAG_PARA_TEXT & 0x3FF) | (1 << 10) | (len(payload) << 20)
        out += struct.pack("<I", header) + payload
    return out


def test_hwp_records_and_control_characters():
    section = _hwp_section(["【청구항 1】", "베이스￼; 및 홀더 본체.", "탭\t뒤"])
    texts = [X.hwp_para_text(p) for tag, _, p in X.hwp_records(section) if tag == X._HWPTAG_PARA_TEXT]
    assert texts == ["【청구항 1】\n", "베이스; 및 홀더 본체.\n", "탭\t뒤\n"]


def test_hwp_ole_extraction_with_fake_ole(monkeypatch):
    section = zlib.compress(_hwp_section(["【청구항 1】", "본문"]))[2:-4]   # raw deflate (no zlib header/trailer)
    header = b"HWP Document File" + b"\0" * (36 - 17) + struct.pack("<I", 1) + b"\0" * 216
    streams = {"FileHeader": header, "BodyText/Section0": section}

    class FakeOle:
        def exists(self, name):
            return name in streams

        def openstream(self, name):
            return io.BytesIO(streams[name])

        def listdir(self):
            return [["FileHeader"], ["BodyText", "Section0"]]

        def close(self):
            pass

    monkeypatch.setattr(X, "olefile", SimpleNamespace(isOleFile=lambda f: True, OleFileIO=lambda f: FakeOle()), raising=False)
    monkeypatch.setitem(__import__("sys").modules, "olefile", SimpleNamespace(isOleFile=lambda f: True, OleFileIO=lambda f: FakeOle()))
    r = X.extract_hwp(b"\xd0\xcf\x11\xe0fake")
    assert r.format == "hwp" and r.text == "【청구항 1】\n본문"
    streams["FileHeader"] = header[:36] + struct.pack("<I", 3) + header[40:]   # encrypted flag
    with pytest.raises(X.ExtractionError, match="암호화"):
        X.extract_hwp(b"\xd0\xcf\x11\xe0fake")


def test_intake_treats_documents_as_text_materials(tmp_path):
    p = tmp_path / "발명설명.hwpx"
    p.write_bytes(_hwpx(["기둥이 판에 구비된다.", "캡을 정렬한다."]))
    item = validate_attachment(p)
    assert item.category == "invention"
    b = MaterialBundle.load(RunRequest(request_text="r", invention_sources=[str(p)]))
    m = b.items[0]
    assert m.kind == "text" and "기둥이 판에 구비된다." in m.text and m.extraction["format"] == "hwpx" and m.as_meta()["extraction"]["chars"] > 0
    bad = tmp_path / "scan.pdf"
    bad.write_bytes(_pdf_with_text(["x"]))
    with pytest.raises(ValueError, match="스캔본"):
        MaterialBundle.load(RunRequest(request_text="r", invention_sources=[str(bad)]))
    assert X.read_text_any(Path(p)).startswith("기둥이")
