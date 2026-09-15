"""Export a run's claims (and optional appendix) as Markdown or DOCX.

The exported claim text is the exact locked wording; the file header records the lock type, revision
identifiers and the exact-text sha256 so a reader can tie the document back to the run records.
"""
from __future__ import annotations

import time
from pathlib import Path

from ..models.state import RunState
from ..pipeline.claimtext import exact_sha256, flatten
from ..pipeline.locks import FINAL_LABEL, PROVISIONAL_LABEL
from .report import _gate_table, render_performance

EAST_ASIAN_FONT = "맑은 고딕"


def _header_rows(state: RunState) -> list[tuple[str, str]]:
    c = state.candidate
    cur = c.current
    dep = state.dependent
    lock = c.final_claim_lock or c.draft_claim_lock
    rows = [
        ("run_id", state.run_id),
        ("상태", (FINAL_LABEL if c.final_claim_lock else PROVISIONAL_LABEL) if lock else "LOCK 없음 — 미확정 문언"),
        ("독립항 LOCK", lock or "없음"),
        ("candidate / revision / design", f"{c.candidate_id} / {cur.revision if cur else c.revision} / {cur.design_revision if cur else c.design_revision}"),
        ("독립항 exact sha256", (cur.exact_sha256 or "") if cur else ""),
    ]
    if dep and dep.current:
        rows += [
            ("종속항 LOCK", dep.final_set_lock or dep.draft_set_lock or "없음"),
            ("dependent_set / dd / dr", f"{dep.dependent_set_id} / {dep.current.dependent_design_revision} / {dep.current.dependent_revision}"),
            ("종속항 세트 sha256", dep.current.exact_sha256 or ""),
        ]
    rows += [("source_set_id", state.source_set_id), ("내보낸 시각", time.strftime("%Y-%m-%d %H:%M"))]
    return rows


def claim_texts(state: RunState) -> tuple[str | None, str | None]:
    cur = state.candidate.current
    root = cur.exact_text if cur and cur.exact_text else None
    dep = state.dependent.current.exact_text if state.dependent and state.dependent.current and state.dependent.current.exact_text else None
    return root, dep


def _unverified(state: RunState) -> list[str]:
    out = []
    if not state.spec_present:
        out.append("OA_FINAL_GATE / DEPENDENT_OA_FINAL_GATE: UNVERIFIED — SPEC_NOT_PROVIDED")
    if not state.prior_art_present:
        out.append("신규성·진보성: UNVERIFIED — PRIOR_ART_NOT_PROVIDED")
    return out


def _evidence_rows(store, state: RunState) -> list[dict]:
    """limitation_evidence rows from the latest design and success records (empty when roles did not fill them)."""
    rows: list[dict] = []
    for kind in ("design", "success", "dependent_success"):
        rid = None
        for r in reversed(list(state.records.values())):
            if r.kind == kind and not r.superseded:
                rid = r.record_id
                break
        if not rid:
            continue
        try:
            rec = store.read_record(state.run_id, rid)
        except FileNotFoundError:
            continue
        for e in rec.get("limitation_evidence") or []:
            rows.append({"record": rid, **e})
    return rows


def render_markdown(store, state: RunState, with_evidence: bool = False) -> str:
    root, dep = claim_texts(state)
    out = ["# 청구항 내보내기", ""]
    for k, v in _header_rows(state):
        out.append(f"- {k}: {v}")
    out += ["", "## 청구범위", ""]
    out.append(root or "독립항 문언 없음")
    if dep:
        out += ["", dep]
    if with_evidence:
        out += ["", "## 부록: 판정 기록", "", _gate_table(state), ""]
        perf = render_performance(state)
        if perf:
            out.append(perf)
        rows = _evidence_rows(store, state)
        if rows:
            out += ["## 부록: 한정별 근거표", "", "| 기록 | 한정 | 역할 | 근거 자료 | 위치 | 근거 종류 | 비고 |", "|---|---|---|---|---|---|---|"]
            for e in rows:
                out.append(f"| {e.get('record', '')} | {e.get('limitation', '')} | {e.get('role', '')} | {e.get('source_name', '')} | {e.get('location', '')} | {e.get('basis', '')} | {e.get('note', '') or ''} |")
            out.append("")
        for u in _unverified(state):
            out.append(f"- {u}")
    return "\n".join(out).rstrip() + "\n"


def _set_font(run, size_pt: float = 11.0) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = EAST_ASIAN_FONT
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), EAST_ASIAN_FONT)


def render_docx(store, state: RunState, path: Path, with_evidence: bool = False) -> Path:
    import docx

    root, dep = claim_texts(state)
    document = docx.Document()
    document.core_properties.title = f"청구항 — {state.run_id}"
    document.core_properties.comments = f"exact_sha256={state.candidate.current.exact_sha256 if state.candidate.current else ''}"
    h = document.add_heading("청구범위", level=1)
    for r in h.runs:
        _set_font(r, 16)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in _header_rows(state):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, str(v)
        for cell in cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    _set_font(r, 9)
    document.add_paragraph("")
    for block in [t for t in (root, dep) if t]:
        for line in block.splitlines():
            p = document.add_paragraph()
            run = p.add_run(line)
            _set_font(run, 11)
            if line.startswith("【청구항"):
                run.bold = True
                p.paragraph_format.space_before = docx.shared.Pt(10)
    if with_evidence:
        h2 = document.add_heading("부록: 판정 기록", level=2)
        for r in h2.runs:
            _set_font(r, 13)
        for line in _gate_table(state).splitlines():
            if line.startswith("|---"):
                continue
            p = document.add_paragraph(line.strip("|").replace("|", " · ").replace("`", ""))
            for r in p.runs:
                _set_font(r, 9)
        rows = _evidence_rows(store, state)
        if rows:
            h3 = document.add_heading("부록: 한정별 근거표", level=2)
            for r in h3.runs:
                _set_font(r, 13)
            t2 = document.add_table(rows=1, cols=6)
            t2.style = "Table Grid"
            for i, name in enumerate(("한정", "역할", "근거 자료", "위치", "근거 종류", "비고")):
                t2.rows[0].cells[i].text = name
            for e in rows:
                cells = t2.add_row().cells
                for i, key in enumerate(("limitation", "role", "source_name", "location", "basis", "note")):
                    cells[i].text = str(e.get(key) or "")
        for u in _unverified(state):
            p = document.add_paragraph(u)
            for r in p.runs:
                _set_font(r, 9)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def verify_docx(path: Path, state: RunState) -> bool:
    """Re-read the exported file and confirm the locked wording survived (flattened comparison)."""
    from ..sources.extract import extract_docx

    text = flatten(extract_docx(path.read_bytes()).text)
    root, dep = claim_texts(state)
    return all(flatten(t) in text for t in (root, dep) if t)


def export_run(store, state: RunState, fmt: str, out: Path | None = None, with_evidence: bool = False) -> Path:
    fmt = fmt.lower()
    if fmt not in ("md", "docx"):
        raise ValueError("지원 형식: md, docx")
    cur = state.candidate.current
    stamp = (cur.revision if cur else state.candidate.revision)
    target = out or (store.run_dir(state.run_id) / f"claims-{state.run_id}-{stamp}.{fmt}")
    if fmt == "md":
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(render_markdown(store, state, with_evidence), encoding="utf-8")
        return target
    render_docx(store, state, target, with_evidence)
    if not verify_docx(target, state):
        raise RuntimeError("DOCX 재검증 실패: 내보낸 문서에서 잠긴 문언을 그대로 읽어내지 못했다")
    return target


def exact_text_hash(state: RunState) -> str | None:
    root, _ = claim_texts(state)
    return exact_sha256(root) if root else None
